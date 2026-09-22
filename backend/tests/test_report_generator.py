"""Informes PDF/DOCX a partir de `estimations.result_json`.

El PDF se genera con WeasyPrint. Si el entorno no tiene Pango/Cairo, los tests
de documento inyectan `mock_pdf_renderer` (marcado explícitamente en los bytes
con `MOCK WeasyPrint`); no se omiten en silencio.
"""

import time
from io import BytesIO
from unittest.mock import MagicMock
from uuid import UUID, uuid4

import pytest
from docx import Document

from app.core.config import settings
from app.models.agent import AgentRunTrace, ToolCallRecord
from app.models.linear_system import LinearSystemInput
from app.services.agent import (
    TOOL_DIAGNOSTICAR_SISTEMA,
    TOOL_RESOLVER_GAUSS,
    TOOL_RESOLVER_GAUSS_JORDAN,
    TOOL_RESOLVER_MATRIZ_INVERSA,
)
from app.services.linear_systems_engine import (
    check_feasibility,
    cross_validate_methods,
    solve_gauss,
    solve_gauss_jordan,
    solve_matrix_inverse,
    validate_system,
    verify_solution,
)
from app.services.markdown_blocks import Paragraph, Table, parse_markdown, spans_text
from app.services.report_generator import (
    EstimationIncompleteError,
    EstimationNotFoundError,
    WeasyPrintUnavailableError,
    build_report_context,
    generate_reports,
    mock_pdf_renderer,
    render_docx,
    render_html,
    render_pdf,
)
from tests.fixtures import panaderia, techchip

A_SMALL = [[2.0, 1.0], [1.0, 3.0]]
B_SMALL = [4.0, 5.0]
A_SINGULAR = [[1.0, 2.0], [2.0, 4.0]]
B_SINGULAR = [1.0, 2.0]


def test_settings_declare_report_storage_fields() -> None:
    assert settings.reports_bucket == "reports"
    assert settings.reports_signed_url_ttl_seconds == 604800


def _estimation_row(
    *,
    A: list[list[float]],
    B: list[float],
    final_response: str,
    variable_names: list[str] | None = None,
    project_name: str = "Planta Norte",
    solvers: bool | None = None,
) -> dict:
    system = LinearSystemInput(A=A, B=B)
    diagnosis = validate_system(system)
    tools = [
        ToolCallRecord(
            name=TOOL_DIAGNOSTICAR_SISTEMA,
            input={"A": A, "B": B},
            output=diagnosis.model_dump(mode="json"),
        )
    ]
    methods = []
    run_solvers = (not diagnosis.is_singular) if solvers is None else solvers
    if run_solvers and not diagnosis.is_singular:
        for solver, tool_name in (
            (solve_gauss, TOOL_RESOLVER_GAUSS),
            (solve_gauss_jordan, TOOL_RESOLVER_GAUSS_JORDAN),
            (solve_matrix_inverse, TOOL_RESOLVER_MATRIZ_INVERSA),
        ):
            solution = solver(system)
            methods.append(solution)
            tools.append(
                ToolCallRecord(
                    name=tool_name,
                    input={"A": A, "B": B},
                    output={
                        "metodo": solution.method.value,
                        "diagnostico": diagnosis.model_dump(mode="json"),
                        **solution.model_dump(mode="json", exclude={"method"}),
                    },
                )
            )
        cross = cross_validate_methods(system, methods)
        substitution = verify_solution(system, cross.solution)
        feasibility = check_feasibility(cross.solution, variable_names)
    else:
        cross = None
        substitution = None
        feasibility = None

    trace = AgentRunTrace(
        A=A,
        B=B,
        variable_names=variable_names,
        tools=tools,
        cross_validation=cross,
        substitution=substitution,
        feasibility=feasibility,
        final_response=final_response,
        model="claude-sonnet-5",
    )
    return {
        "id": str(uuid4()),
        "problem_text": "¿Qué plan de producción cabe en la planta esta semana?",
        "result_json": trace.model_dump(mode="json"),
        "projects": {"name": project_name},
    }


def test_context_includes_every_gauss_step_not_just_x() -> None:
    row = _estimation_row(
        A=A_SMALL,
        B=B_SMALL,
        final_response="Plan viable: producir 1.4 y 1.2 miles de módulos.",
        variable_names=["AI-Edge 1", "AI-Server Pro"],
    )
    ctx = build_report_context(row)
    gauss = ctx.methods[0]
    assert gauss.ran is True
    assert gauss.step_blocks
    assert any("F" in step.description for step in gauss.step_blocks)
    assert ctx.methods_agree is True
    assert "coinciden" in ctx.comparison_verdict
    assert "plan de asignación" in " ".join(ctx.conclusion_paragraphs).lower()
    html = render_html(ctx)
    assert "Resolución multimétodo" in html
    assert "Gauss-Jordan" in html
    assert "Matriz inversa" in html
    assert "Resumen ejecutivo" in html
    assert "Plan viable" in html

    document = Document(BytesIO(render_docx(ctx)))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "Planta Norte" in texts
    assert any("Resolución multimétodo" == text for text in texts)
    assert any("Diagnóstico y conclusiones" == text for text in texts)
    assert any("E = |AX − B|" in text or "E = |AX" in text for text in texts)


def test_singular_report_explains_business_alert_without_x() -> None:
    row = _estimation_row(
        A=A_SINGULAR,
        B=B_SINGULAR,
        final_response="No hay un plan único: las restricciones son linealmente dependientes.",
    )
    ctx = build_report_context(row)
    assert all(not method.ran for method in ctx.methods)
    assert ctx.needs_alert is True
    conclusiones = " ".join(ctx.conclusion_paragraphs).lower()
    assert "singular" in conclusiones
    # El informe no asume manufactura: sirve igual para reparto, asignación o mezcla.
    assert "plan de asignación" in conclusiones
    assert "solución del sistema" in conclusiones
    assert "orden de fabricación" not in conclusiones
    assert "plan de producción" not in conclusiones
    html = render_html(ctx)
    assert "Sistema singular" in html or "singular" in html.lower()
    assert "No hay tabla comparativa" in html
    assert "Estimación por sistemas de ecuaciones lineales" in html
    assert "Estimación de plan de producción" not in html


def test_infeasible_flag_is_translated_to_business_language() -> None:
    row = _estimation_row(
        A=techchip.A,
        B=techchip.B_LITERAL_GUIA,
        final_response="Hay un X matemático, pero no se puede fabricar.",
        variable_names=techchip.VARIABLE_NAMES,
    )
    ctx = build_report_context(row)
    assert ctx.needs_alert is True
    joined = " ".join(ctx.conclusion_paragraphs).lower()
    assert "componente de solución negativo" in joined
    assert "negativ" in joined
    html = render_html(ctx)
    assert "infactibilidad" in html.lower() or "negativo" in html.lower()


def test_missing_result_json_is_incomplete() -> None:
    with pytest.raises(EstimationIncompleteError):
        build_report_context(
            {
                "id": str(uuid4()),
                "problem_text": "sin resultado",
                "result_json": None,
                "projects": {"name": "X"},
            }
        )


def test_render_pdf_with_weasyprint_or_documented_mock() -> None:
    html = "<html><body><h1>Informe</h1></body></html>"
    try:
        pdf = render_pdf(html)
    except WeasyPrintUnavailableError:
        pdf = mock_pdf_renderer(html)
        assert b"MOCK WeasyPrint" in pdf
        return
    assert pdf.startswith(b"%PDF")
    assert b"MOCK WeasyPrint" not in pdf


def test_render_pdf_raises_when_pango_is_missing(monkeypatch: pytest.MonkeyPatch) -> None:
    def boom(html: str, base_url: str) -> bytes:
        raise OSError("cannot load library 'pango-1.0-0'")

    monkeypatch.setattr(
        "app.services.report_generator._weasyprint_write_pdf", boom
    )
    with pytest.raises(WeasyPrintUnavailableError, match="Pango"):
        render_pdf("<html><body>x</body></html>")


def test_generate_reports_uploads_both_files_using_documented_pdf_mock() -> None:
    row = _estimation_row(
        A=A_SMALL,
        B=B_SMALL,
        final_response="Resumen listo para gerencia.",
        variable_names=["Línea A", "Línea B"],
    )
    estimation_id = UUID(row["id"])
    client = MagicMock()
    select = (
        client.table.return_value.select.return_value.eq.return_value.limit.return_value
    )
    select.execute.return_value = MagicMock(data=[row])

    generated_at = "2026-09-19T04:00:00+00:00"
    client.table.return_value.insert.return_value.execute.return_value = MagicMock(
        data=[
            {
                "id": str(uuid4()),
                "file_url": "https://signed.example/informe.docx",
                "file_type": "docx",
                "generated_at": generated_at,
            },
            {
                "id": str(uuid4()),
                "file_url": "https://signed.example/informe.pdf",
                "file_type": "pdf",
                "generated_at": generated_at,
            },
        ]
    )
    bucket = client.storage.from_.return_value

    def _signed(path: str, _ttl: int, options=None):
        name = str(path).rsplit("/", 1)[-1]
        url = f"https://signed.example/{name}"
        return {"signedURL": url, "signedUrl": url}

    bucket.create_signed_url.side_effect = _signed

    response = generate_reports(
        estimation_id, client, pdf_renderer=mock_pdf_renderer
    )

    assert response.estimation_id == estimation_id
    assert [item.file_type for item in response.reports] == ["docx", "pdf"]
    assert {item.file_url for item in response.reports} == {
        "https://signed.example/informe.docx",
        "https://signed.example/informe.pdf",
    }
    # Una firma al persistir cada archivo y otra al responder (URLs de descarga frescas).
    assert bucket.create_signed_url.call_count == 4
    assert bucket.upload.call_count == 2
    mime_types = [
        (call.kwargs.get("file_options") or {})["content-type"]
        for call in bucket.upload.call_args_list
    ]
    assert "application/pdf" in mime_types
    assert (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        in mime_types
    )
    assert all(isinstance(call.args[1], (bytes, bytearray)) for call in bucket.upload.call_args_list)
    assert any(b"%PDF" in call.args[1] for call in bucket.upload.call_args_list)
    assert any(call.args[1].startswith(b"PK") for call in bucket.upload.call_args_list)


RESUMEN_MARKDOWN = """\
## Diagnóstico

El sistema es **compatible determinado**. det(A) = -10.

## Plan de lotes

| Pan | Lotes |
| --- | --- |
| pan de caja | 4 |
| baguette | 3 |

```
x = (4, 3, 2)
```

La operación (0.5)*F1 no es una cursiva.
"""


def _executive_slice(document: Document) -> tuple[list[str], list]:
    """Párrafos y tablas que pertenecen solo al resumen, antes de los métodos."""
    texts: list[str] = []
    started = False
    for paragraph in document.paragraphs:
        if paragraph.text == "Resumen ejecutivo":
            started = True
            continue
        if paragraph.text == "Resolución multimétodo":
            break
        if started:
            texts.append(paragraph.text)
    # La primera tabla del documento es la del resumen; las de los pasos vienen después.
    return texts, list(document.tables[:1])


def test_el_resumen_markdown_no_se_imprime_con_simbolos_literales() -> None:
    row = _estimation_row(
        A=panaderia.A,
        B=panaderia.B,
        final_response=RESUMEN_MARKDOWN,
        variable_names=panaderia.VARIABLE_NAMES,
    )
    ctx = build_report_context(row)
    html = render_html(ctx)
    resumen = html.split("<h2>Resolución multimétodo</h2>", 1)[0]

    assert "<h4>Diagnóstico</h4>" in resumen
    assert "<strong>compatible determinado</strong>" in resumen
    assert "<th" in resumen and "pan de caja" in resumen
    assert "<pre>" in resumen and "x = (4, 3, 2)" in resumen
    assert "##" not in resumen
    assert "**" not in resumen
    assert "|---|" not in resumen
    assert "(0.5)*F1" in resumen

    document = Document(BytesIO(render_docx(ctx)))
    texts, tables = _executive_slice(document)
    joined = "\n".join(texts)
    assert "Diagnóstico" in joined
    assert "##" not in joined
    assert "**" not in joined
    assert "|---|" not in joined
    assert "(0.5)*F1" in joined
    assert tables and tables[0].rows[1].cells[0].text == "pan de caja"

    bold = [
        run.text
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.bold and "compatible determinado" in run.text
    ]
    assert bold


def test_panaderia_gauss_jordan_y_la_inversa_no_comparten_la_matriz() -> None:
    """Las operaciones de fila coinciden (solo dependen de A); el estado, no."""
    row = _estimation_row(
        A=panaderia.A,
        B=panaderia.B,
        final_response=RESUMEN_MARKDOWN,
        variable_names=panaderia.VARIABLE_NAMES,
        project_name="Panadería",
    )
    ctx = build_report_context(row)
    jordan = ctx.methods[1]
    inversa = ctx.methods[2]

    assert jordan.ran and inversa.ran
    assert [step.description for step in jordan.step_blocks] == [
        step.description for step in inversa.step_blocks
    ]
    assert jordan.step_blocks and inversa.step_blocks
    assert all(len(step.matrix[0]) == 4 for step in jordan.step_blocks)
    assert all(len(step.matrix[0]) == 6 for step in inversa.step_blocks)
    assert jordan.inverse_matrix is None
    assert inversa.inverse_matrix and len(inversa.inverse_matrix[0]) == 3
    assert not jordan.component_steps
    assert inversa.component_steps
    assert 'class="bar"' in jordan.step_blocks[0].matrix_html
    assert 'class="bar"' in inversa.step_blocks[0].matrix_html

    html = render_html(ctx)
    jordan_html, inversa_html = html.split("<h3>Matriz inversa</h3>", 1)
    jordan_html = jordan_html.split("<h3>Gauss-Jordan</h3>", 1)[1]
    assert jordan_html.count("<td") > 0
    # Una fila de [A|B] tiene 4 celdas; una de [A|I], 6.
    assert "<td" in jordan_html and jordan_html.count("<td") % 4 == 0
    inversa_steps = inversa_html.split("Matriz inversa A", 1)[0]
    assert inversa_steps.count("<td") % 6 == 0

    document = Document(BytesIO(render_docx(ctx)))
    widths = [len(table.columns) for table in document.tables]
    assert 4 in widths and 6 in widths
    # La barra de ampliación es un borde izquierdo más grueso (sz 16) en la columna de B o de I.
    barred = 0
    for table in document.tables:
        if len(table.columns) not in (4, 6):
            continue
        for cell in table.rows[0].cells:
            borders = cell._tc.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr"
            )
            if borders is not None and 'w:sz="16"' in borders.xml:
                barred += 1
    assert barred > 0


def _parse_within_a_second(text: str) -> list:
    started = time.perf_counter()
    blocks = parse_markdown(text)
    assert time.perf_counter() - started < 1.0
    return blocks


def _joined(blocks: list) -> str:
    parts: list[str] = []
    for block in blocks:
        if isinstance(block, Paragraph):
            parts.append(spans_text(block.spans))
        elif isinstance(block, Table):
            parts.append(spans_text(block.header[0]) if block.header else "")
    return "\n".join(parts)


def test_parse_frase_con_barras_de_matriz_no_se_cuelga() -> None:
    """La prosa que colgó el informe: [A|B] y [I|X] no son una tabla."""
    text = (
        "Reduce [A|B] hasta [I|X] con operaciones de escalado y eliminación "
        "en ambas direcciones, llegando directamente a:"
    )
    blocks = _parse_within_a_second(text)
    assert not any(isinstance(block, Table) for block in blocks)
    assert "Reduce [A|B] hasta [I|X]" in _joined(blocks)


def test_parse_multiples_barras_sueltas_no_forman_tabla() -> None:
    text = "harina | horno | amasado | sobrante | sin separador de tabla"
    blocks = _parse_within_a_second(text)
    assert not any(isinstance(block, Table) for block in blocks)
    assert "harina | horno | amasado" in _joined(blocks)


def test_parse_tabla_seguida_de_barra_suelta_no_absorbe_la_prosa() -> None:
    text = """\
| Pan | Lotes |
| --- | --- |
| pan de caja | 4 |
Reduce [A|B] hasta [I|X] y esta línea no es una fila.
"""
    blocks = _parse_within_a_second(text)
    tables = [block for block in blocks if isinstance(block, Table)]
    assert len(tables) == 1
    assert len(tables[0].rows) == 1
    assert spans_text(tables[0].rows[0][0]) == "pan de caja"
    assert any(
        isinstance(block, Paragraph) and "[A|B]" in spans_text(block.spans)
        for block in blocks
    )


def test_parse_linea_con_barras_en_los_extremos_sin_gfm() -> None:
    """Empieza y termina en `|`, pero no hay fila `---` que la vuelva tabla."""
    text = "|esto no es | una tabla|"
    blocks = _parse_within_a_second(text)
    assert not any(isinstance(block, Table) for block in blocks)
    assert "|esto no es | una tabla|" in _joined(blocks)


def test_generate_reports_404_when_estimation_missing() -> None:
    client = MagicMock()
    select = (
        client.table.return_value.select.return_value.eq.return_value.limit.return_value
    )
    select.execute.return_value = MagicMock(data=[])
    with pytest.raises(EstimationNotFoundError):
        generate_reports(uuid4(), client, pdf_renderer=mock_pdf_renderer)
